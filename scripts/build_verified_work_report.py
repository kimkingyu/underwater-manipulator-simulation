"""Read-only audit of existing project data; render the evidence-based report.
Does not run MATLAB/Fluent, modify source data, or overwrite earlier reports.
"""
from pathlib import Path
import csv
import json
import hashlib
import collections
import subprocess
import math
import re
import sys

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs'
STEM = '水下机械臂_前期工作详细报告_证据核验版'


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def audit():
    matrix = ROOT / 'cfd/parametric_cfd_matrix_52'
    summary_path = matrix / 'cfd_52_matrix_summary.json'
    summary = {x['case_id']: x for x in json.loads(summary_path.read_text(encoding='utf-8-sig'))}
    cases = []
    for folder in sorted(matrix.glob('case_*')):
        p = folder / 'moments.csv'
        sc_path = folder / 'moments.validation.json'
        rows = list(csv.DictReader(p.open(encoding='utf-8-sig', newline='')))
        sc = json.loads(sc_path.read_text(encoding='utf-8-sig'))
        torques = [[float(r[f'j{j}_torque_nm']) for j in (1, 2, 3)] for r in rows]
        item = {'case_id': folder.name, 'rows': len(rows), 'time_end_s': max(float(r['flow_time_s']) for r in rows),
                'csv_sha256': sha(p), 'sidecar_sha256': sha(sc_path), 'valid_flag': sc.get('valid'),
                'status_flag': sc.get('status'), 'selected_step': summary[folder.name]['step_used'],
                'max_abs_recorded_torque_nm': max(abs(v) for t in torques for v in t),
                'recorded_step_seconds': [float(r['step_elapsed_s']) for r in rows if r.get('step_elapsed_s')],
                'summary_matches_selected_csv': None}
        selected = [r for r in rows if int(r['step_index']) == item['selected_step']]
        if selected:
            diffs = [float(summary[folder.name][f'j{j}_cfd_nm']) - float(selected[0][f'j{j}_torque_nm']) for j in (1, 2, 3)]
            item['summary_minus_csv_nm'] = diffs
            item['summary_matches_selected_csv'] = all(abs(v) < 1e-8 for v in diffs)
        cases.append(item)
    duplicate_groups = collections.defaultdict(list)
    for c in cases:
        duplicate_groups[c['csv_sha256']].append(c['case_id'])
    historical_matches = []
    historical = ROOT / 'cfd/parametric_cfd_matrix'
    for old in historical.glob('case_*/moments.csv'):
        old_hash = sha(old)
        for c in cases:
            if old_hash == c['csv_sha256']:
                historical_matches.append({'matrix52': c['case_id'], 'historical9': old.parent.name})
    duration = [v for c in cases for v in c['recorded_step_seconds'] if math.isfinite(v)]
    report_path = ROOT / 'cfd/wetted_arm_rebuild/cfd_vs_matlab_report.json'
    report = json.loads(report_path.read_text(encoding='utf-8'))
    short_metrics = []
    for j in range(3):
        errors = [b[j] - a[j] for a, b in zip(report['cfd_tau'], report['matlab_tau'])]
        short_metrics.append({'joint': j+1, 'mae_nm': sum(abs(e) for e in errors)/len(errors),
                              'rmse_nm': math.sqrt(sum(e*e for e in errors)/len(errors)),
                              'stored_r2': report['metrics'][j]['R2']})
    matlab_extract = r'''
import json, scipy.io as s, numpy as n, h5py
m=s.loadmat('data/hydro_surrogate/full/model.mat',struct_as_record=False,squeeze_me=True)['model']
d=s.loadmat('data/joint_loads_data.mat',struct_as_record=False,squeeze_me=True)
c=s.loadmat('data/underarm_3dof_sim_data.mat',struct_as_record=False,squeeze_me=True)
r=s.loadmat('data/hydro_surrogate/full/report.mat',struct_as_record=False,squeeze_me=True)['report']
res={'model_kind':str(m.kind),'model_val_mae_nm':n.asarray(m.val_mae).tolist(),'model_test_mae_nm':n.asarray(m.test_mae).tolist(), 'stored_report_note':str(r.note), 'stored_audit_note':str(r.auditNote),
'load_time_s':[float(d['time'][0]),float(d['time'][-1])], 'load_samples':len(d['time']), 'load_peak_total_nm':n.max(n.abs(d['tau_total']),axis=0).tolist(), 'load_rms_total_nm':n.sqrt(n.mean(d['tau_total']**2,axis=0)).tolist(), 'load_peak_drag_nm':n.max(n.abs(d['tau_hydro_drag']),axis=0).tolist(), 'current_speed_mps':float(d['env'].current_speed),'current_azimuth_rad':float(d['env'].current_psi), 'base_link3_min_clearance_m':float(n.min(d['clearance_base_link3'])), 'link2_link4_nan_samples':int(n.sum(n.isnan(d['clearance_link2_link4']))), 'tcp_extent_m':n.ptp(d['P_tcp_des'],axis=0).tolist(), 'absolute_mechanical_work_J':float(n.trapz(n.sum(d['joint_mech_power'],axis=1),d['time'])), 'net_mechanical_work_J':float(n.trapz(n.sum(d['tau_total']*d['qd_seq'],axis=1),d['time'])), 'control_samples':len(c['time']), 'control_time_s':[float(c['time'][0]),float(c['time'][-1])], 'control_max_joint_error_deg':(n.max(n.abs(c['log_q']-c['q_des']),axis=0)*180/n.pi).tolist(), 'control_torque_limit_assumption_nm':float(c['max_torque_limit'])}
with h5py.File('cfd/wetted_arm_rebuild/articulated_mesh_v2/cad_attempt04/reduced_412k_system/assembled_reduced_415k_t0.cas.h5','r') as f:
    g=f['meshes/1/cells/zoneTopology']
    res['mesh_cell_zone_ids']=g['id'][()].tolist()
    res['mesh_zone_cell_counts']=(g['maxId'][()]-g['minId'][()]+1).tolist()
    res['mesh_total_cells']=int(n.sum(g['maxId'][()]-g['minId'][()]+1))
print(json.dumps(res))
'''
    proc = subprocess.run([r'D:\AMESIM\Amesim\sys\python\win64\python.exe', '-c', matlab_extract], cwd=ROOT, capture_output=True, text=True, check=True)
    mat = json.loads(proc.stdout)
    sources = [summary_path, report_path, ROOT/'src/cfd/cfd_case_parameters.m', ROOT/'src/dynamics/underwater_hydro_defaults.m',
               ROOT/'scripts/generate_full_comparison_excel.py', ROOT/'cfd/verify_and_clean_52_matrix.py',
               ROOT/'data/joint_loads_data.mat',ROOT/'data/underarm_3dof_sim_data.mat',ROOT/'data/hydro_surrogate/full/model.mat',
               ROOT/'data/hydro_surrogate/full/report.mat',ROOT/'fluent-20260924-014845-33960.trn',
               ROOT/'cfd/wetted_arm_rebuild/articulated_mesh_v2/cad_attempt04/reduced_412k_system/assembled_reduced_415k_t0.cas.h5']
    result = {'audit_date':'2026-09-29','scope':'Existing local artifacts only. No new solver run. Hashes identify snapshots, not authenticity certificates.',
              'case_count':len(cases), 'row_count_distribution':dict(collections.Counter(c['rows'] for c in cases)),
              'selected_step_distribution':dict(collections.Counter(c['selected_step'] for c in cases)),
              'duplicate_csv_groups':[x for x in duplicate_groups.values() if len(x)>1], 'historical_csv_matches':historical_matches,
              'summary_csv_mismatches':[c['case_id'] for c in cases if not c['summary_matches_selected_csv']],
              'recorded_step_seconds_summary':{'n':len(duration),'min':min(duration),'max':max(duration),'mean':sum(duration)/len(duration)},
              'short_30_step_metrics':short_metrics,'matlab_archive':mat,'cases':cases,
              'source_sha256':{str(p.relative_to(ROOT)).replace(chr(92),'/'):sha(p) for p in sources}}
    path = OUT/(STEM+'_核验记录.json')
    path.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({k:v for k,v in result.items() if k not in ('cases','source_sha256')},ensure_ascii=False,indent=2))


def render():
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    source = OUT/(STEM+'.md')
    text = source.read_text(encoding='utf-8')
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21)
    sec.top_margin, sec.bottom_margin = Cm(2.4), Cm(2.2)
    sec.left_margin, sec.right_margin = Cm(2.5), Cm(2.3)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
        style = doc.styles[name]
        style.font.name = 'Times New Roman'
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体' if name == 'Normal' else '黑体')
        style.font.color.rgb = RGBColor.from_string('18324D' if name != 'Normal' else '222222')
    normal = doc.styles['Normal']
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(6)
    doc.styles['Title'].font.size = Pt(22)
    doc.styles['Heading 1'].font.size = Pt(15)
    doc.styles['Heading 2'].font.size = Pt(12)
    doc.styles['Heading 3'].font.size = Pt(11)
    header = sec.header.paragraphs[0]
    header.text = '水下机械臂项目  |  前期工作详细报告（证据核验版）'
    header.runs[0].font.size = Pt(9)
    footer = sec.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run('第 ')
    f = OxmlElement('w:fldSimple'); f.set(qn('w:instr'),'PAGE'); footer._p.append(f)
    footer.add_run(' 页')
    def clean(t):
        t = re.sub(r'\[([^\]]+)\]\(([^)]+)\)',r'\1',t)
        return t.replace('**','').replace('`','')
    lines = text.splitlines(); i = 0
    tables = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip() or line.strip() == '---':
            i += 1; continue
        if line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].startswith('|'):
                raw=[clean(x.strip()) for x in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'[:\- ]+',x or ' ') for x in raw): rows.append(raw)
                i+=1
            cols=max(map(len,rows)); table=doc.add_table(rows=1,cols=cols)
            table.style='Light Shading Accent 1'
            for j,v in enumerate(rows[0]):table.rows[0].cells[j].text=v
            repeat=OxmlElement('w:tblHeader'); table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for row in rows[1:]:
                cells=table.add_row().cells
                for j,v in enumerate(row): cells[j].text=v
            for row in table.rows:
                cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
                for cell in row.cells:
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after=Pt(3)
                        for run in p.runs:
                            run.font.size=Pt(9)
                            run.font.name='Times New Roman'
                            run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
            doc.add_paragraph(); tables+=1; continue
        if line.startswith('# '):
            doc.add_paragraph(clean(line[2:]),'Title')
        elif line.startswith('## '):
            doc.add_heading(clean(line[3:]),1)
        elif line.startswith('### '):
            doc.add_heading(clean(line[4:]),2)
        else:
            doc.add_paragraph(clean(line))
        i+=1
    doc.core_properties.title='水下机械臂前期工作详细报告（证据核验版）'
    doc.core_properties.subject='本地工程文件核查、阶段性工作与验证范围'
    doc.core_properties.author=''
    dest=OUT/(STEM+'.docx')
    doc.save(dest)
    check=Document(dest)
    assert len(check.tables)==tables
    assert len(check.paragraphs)>80
    for p in check.paragraphs:
        assert '\ufffd' not in p.text
    print(json.dumps({'docx':str(dest),'paragraphs':len(check.paragraphs),'tables':len(check.tables),'bytes':dest.stat().st_size},ensure_ascii=False))

if __name__=='__main__':
    if '--render' in sys.argv: render()
    else: audit()
